import streamlit as st
import openpyxl
from openpyxl.utils import get_column_letter
import pandas as pd
import csv
import io
import shutil
import tempfile
import os


st.set_page_config(
    page_title="CSV → Excel コピーツール",
    page_icon="📊",
    layout="centered"
)

st.title("📊 CSV → Excel コピーツール")
st.markdown("---")

# ========== Excelファイル設定 ==========
st.header("① Excelファイル設定")

template_file = st.file_uploader(
    "テンプレートExcelをアップロード",
    type=["xlsx", "xlsm"],
    key="template"
)

output_filename = st.text_input(
    "出力ファイル名",
    value="出力ファイル.xlsx"
)
if not output_filename.endswith(".xlsx"):
    output_filename += ".xlsx"

st.markdown("---")

# ========== 測定条件 ==========
st.header("② 測定条件")

col1, col2 = st.columns(2)

with col1:
    place_options = ["FF garage", "FF parking", "その他"]
    place = st.selectbox("Place", place_options)
    if place == "その他":
        place = st.text_input("Place（入力）", key="place_custom")

    car_options = [
        "FF B33 demo car", "FF B25 demo car", "FE B33 demo car",
        "FE B25 demo car", "FHP B33 demo car", "その他"
    ]
    car = st.selectbox("Car", car_options)
    if car == "その他":
        car = st.text_input("Car（入力）", key="car_custom")

    dsp = st.text_input("DSP data")

    filter_val = st.selectbox("filter", ["bypass", "ON"])

    # gain（+5.00 ～ -30.00、0.25刻み）
    gain_values = []
    v = 5.0
    while v >= -30.0:
        gain_values.append(round(v, 2))
        v = round(v - 0.25, 2)
    gain_idx = gain_values.index(0.0)
    gain = st.selectbox(
        "gain",
        options=gain_values,
        index=gain_idx,
        format_func=lambda x: f"{x:.2f} dB"
    )

with col2:
    eq = st.selectbox("EQ", ["bypass", "ON"])

    signal = st.selectbox("Signal", ["Pink Noise", "Sin wave", "White Noise"])

    mic_pos_options = ["FR Sheet", "FL Sheet", "RR Sheet", "RL Sheet", "その他"]
    mic_pos = st.selectbox("Mic position", mic_pos_options)
    if mic_pos == "その他":
        mic_pos = st.text_input("Mic position（入力）", key="mic_pos_custom")

    mic = st.selectbox("Mic", ["Binaural", "1 point"])

st.markdown("---")

# ========== CSVファイル設定 ==========
st.header("③ CSVファイル設定")

st.subheader("フロアノイズCSV")
floor_csv_file = st.file_uploader(
    "フロアノイズCSVをアップロード（1ファイルのみ）",
    type=["csv"],
    key="floor_csv"
)

st.subheader("測定データCSV")
data_csv_files = st.file_uploader(
    "測定データCSVをアップロード（複数可）",
    type=["csv"],
    accept_multiple_files=True,
    key="data_csv"
)

# 測定データの順番表示
if data_csv_files:
    st.write("**選択されたファイル（この順番でシートが作成されます）:**")
    for i, f in enumerate(data_csv_files):
        st.write(f"　{i+1}. {f.name}")

st.markdown("---")

# ========== 実行 ==========
def read_csv_from_upload(uploaded_file):
    """アップロードされたCSVを読み込む"""
    content = uploaded_file.read().decode("utf-8-sig")
    reader = csv.reader(content.splitlines())
    return list(reader)

def write_range(ws, start_row, start_col, data):
    """2次元リストをExcelに書き込む"""
    if not data or not data[0]:
        return
    for i, row in enumerate(data):
        for j, value in enumerate(row):
            ws.cell(row=start_row + i, column=start_col + j, value=value)

def process_csv(ws, csv_data, csv_filename, place, car,
                dsp, filter_val, gain, eq, signal, mic_pos, mic,
                is_floor_noise=False):
    """CSVデータをExcelシートに書き込む"""
    # 1~3行目を転置して D39 から
    rows_1to3 = csv_data[0:3]
    transposed = list(map(list, zip(*rows_1to3)))
    write_range(ws, 39, 4, transposed)

    # 5行A列 → T2
    if len(csv_data) >= 5:
        ws.cell(row=2, column=20, value=csv_data[4][0])

    # T列
    ws.cell(row=3, column=20, value=place)
    ws.cell(row=4, column=20, value=car)
    ws.cell(row=5, column=20, value=csv_filename)

    # V列（フロアノイズ以外）
    if not is_floor_noise:
        ws.cell(row=2, column=22, value=dsp)
        ws.cell(row=3, column=22, value=filter_val)
        ws.cell(row=4, column=22, value=gain)
        ws.cell(row=5, column=22, value=eq)
        ws.cell(row=2, column=24, value=signal)

    # X列
    ws.cell(row=3, column=24, value=mic_pos)
    ws.cell(row=4, column=24, value=mic)


if st.button("🚀 実行！", use_container_width=True, type="primary"):

    # 入力チェック
    errors = []
    if not template_file:
        errors.append("テンプレートExcelをアップロードしてください")
    if not place:
        errors.append("Placeを入力してください")
    if not car:
        errors.append("Carを入力してください")
    if not mic_pos:
        errors.append("Mic positionを入力してください")
    if not data_csv_files:
        errors.append("測定データCSVをアップロードしてください")
    if not floor_csv_file:
        errors.append("フロアノイズCSVをアップロードしてください")

    if errors:
        for e in errors:
            st.error(e)
    else:
        with st.spinner("⏳ ファイル生成中...しばらくお待ちください"):
            try:
                # テンプレートを読み込む
                template_bytes = template_file.read()
                wb = openpyxl.load_workbook(io.BytesIO(template_bytes))

                # シート確認
                if "テンプレート" not in wb.sheetnames:
                    st.error(f"「テンプレート」シートがありません\n現在のシート: {wb.sheetnames}")
                    st.stop()
                if "フロアノイズ" not in wb.sheetnames:
                    st.error(f"「フロアノイズ」シートがありません\n現在のシート: {wb.sheetnames}")
                    st.stop()

                # フロアノイズをリネーム
                wb["フロアノイズ"].title = "フロアノイズ_元"

                # テンプレートシートを取得
                template_ws = wb["テンプレート"]

                # 測定データCSV処理
                for entry in data_csv_files:
                    sheet_name = os.path.splitext(entry.name)[0][:31]
                    csv_filename = entry.name
                    entry.seek(0)
                    csv_data = read_csv_from_upload(entry)

                    # テンプレートシートをコピー
                    new_ws = wb.copy_worksheet(template_ws)
                    new_ws.title = sheet_name

                    process_csv(new_ws, csv_data, csv_filename,
                                place, car, dsp, filter_val, gain, eq,
                                signal, mic_pos, mic, is_floor_noise=False)

                # フロアノイズCSV処理
                floor_ws_orig = wb["フロアノイズ_元"]
                new_floor_ws = wb.copy_worksheet(floor_ws_orig)
                new_floor_ws.title = "フロアノイズ"
                floor_csv_file.seek(0)
                floor_data = read_csv_from_upload(floor_csv_file)
                floor_filename = floor_csv_file.name

                process_csv(new_floor_ws, floor_data, floor_filename,
                            place, car, dsp, filter_val, gain, eq,
                            signal, mic_pos, mic, is_floor_noise=True)

                # 元のテンプレートシートを削除
                del wb["テンプレート"]
                del wb["フロアノイズ_元"]

                # フロアノイズを最後に移動
                sheets = wb.sheetnames
                floor_idx = sheets.index("フロアノイズ")
                last_idx = len(sheets) - 1
                if floor_idx != last_idx:
                    wb.move_sheet("フロアノイズ", offset=last_idx - floor_idx)

                # バイト列として保存
                output = io.BytesIO()
                wb.save(output)
                output.seek(0)

                st.success("✅ 完了！下のボタンからダウンロードしてください")

                # ダウンロードボタン
                st.download_button(
                    label="📥 Excelファイルをダウンロード",
                    data=output.getvalue(),
                    file_name=output_filename,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )

            except Exception as e:
                import traceback
                st.error(f"エラーが発生しました:\n{traceback.format_exc()}")