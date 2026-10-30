#!/usr/bin/env python3
"""
csv_to_excel_sweep_gh.py — Sine Sweep RTA CSV → Excel (openpyxl only / GitHub版)

依存: openpyxl のみ (win32com 不要)
チャート: ScatterChart + 真の対数X軸 (logBase=10) + log paper グリッド
"""

import math, os, sys, re, zipfile, shutil, tempfile, datetime, json
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from openpyxl import load_workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ── 定数 ─────────────────────────────────────────────────────────────────────
SCRIPT_DIR    = os.path.dirname(os.path.abspath(__file__))
SETTINGS_FILE = os.path.join(SCRIPT_DIR, "settings_sweep_gh.json")

A = "http://schemas.openxmlformats.org/drawingml/2006/main"

LABEL_FREQS = [10, 30, 50, 70, 100, 300, 500, 700,
               1000, 3000, 5000, 7000, 10000, 20000]

COLOR_BG  = "#f5f5f5"
COLOR_SEC = "#e8e8e8"
COLOR_BTN = "#4472c4"
COLOR_RUN = "#2e7d32"
COLOR_FG  = "#ffffff"

PLACE_OPTIONS   = ["FF garage", "FF parking", "その他"]
CAR_OPTIONS     = ["FF B33 demo car", "FF B25 demo car",
                   "FE B33 demo car", "FE B25 demo car",
                   "FHP B33 demo car", "その他"]
SIGNAL_OPTIONS  = ["Sin Sweep", "Pink Noise", "White Noise"]
MIC_POS_OPTIONS = ["FR Sheet", "FL Sheet", "RR Sheet", "RL Sheet", "その他"]
MIC_OPTIONS     = ["Binaural", "1 point"]

DEFAULT_SETTINGS = {
    "place": "FF garage", "place_custom": "",
    "car": "FF B25 demo car", "car_custom": "",
    "signal": "Sin Sweep",
    "mic": "Binaural", "mic_pos": "FR Sheet", "mic_pos_custom": "",
    "note": "", "output_dir": "", "output_name": "出力結果_Sweep",
}

_PALETTE = [
    "4472C4", "ED7D31", "70AD47", "C0504D",
    "5B9BD5", "92D050", "FFC000", "A46BBE",
]


# ── ヘルパー ──────────────────────────────────────────────────────────────────

def freq_to_label(v):
    if v >= 1000:
        k = v / 1000
        return f"{int(k)}k" if k == int(k) else f"{k:g}k"
    return str(int(v)) if v == int(v) else f"{v:g}"


# ── CSV パーサー ───────────────────────────────────────────────────────────────

def parse_sine_csv(path: str) -> dict:
    """Sine Sweep CSV → {freq: spl}"""
    data = {}
    in_data = header_skipped = False
    try:
        with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
            for line in f:
                line = line.strip()
                if line == "[DATA]":
                    in_data = True; continue
                if not in_data: continue
                if line.startswith("COMMENT="): continue
                if not header_skipped:
                    header_skipped = True; continue
                parts = line.split(",")
                if len(parts) < 2: continue
                try:
                    freq, spl = float(parts[0]), float(parts[1])
                    if 10 <= freq <= 25000:
                        data[freq] = spl
                except ValueError:
                    continue
    except Exception:
        pass
    return data


def extract_group_level(filepath: str) -> tuple:
    name  = os.path.splitext(os.path.basename(filepath))[0]
    parts = name.rsplit("_", 1)
    return (parts[0], parts[1]) if len(parts) == 2 else (name, "")


# ── Excel / Chart 生成 ────────────────────────────────────────────────────────

def _thin_border():
    t = Side(border_style="thin", color="000000")
    return Border(left=t, right=t, top=t, bottom=t)


def _cell(ws, r, c, v, bold=False, fill=None, size=9):
    cell = ws.cell(r, c, v)
    cell.font      = Font(name="Meiryo UI", bold=bold, size=size)
    cell.border    = _thin_border()
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    if fill:
        cell.fill = PatternFill(start_color=fill, end_color=fill, fill_type="solid")
    return cell


def _make_sheet(wb, group_name: str, entries: list, settings: dict):
    ws = wb.create_sheet(group_name)
    DATA_START = 40
    FREQ_COL   = 4
    DATA_COL   = 5

    # メタデータ
    date_val = settings.get("date", datetime.date.today().isoformat())
    place    = settings.get("place") or settings.get("place_custom", "")
    car      = settings.get("car")   or settings.get("car_custom", "")
    for i, (lbl, val) in enumerate([
        ("Date",   date_val),
        ("Place",  place),
        ("Car",    car),
        ("Signal", settings.get("signal", "Sin Sweep")),
        ("Mic",    settings.get("mic", "")),
    ]):
        r = 2 + i
        ws.row_dimensions[r].height = 16
        _cell(ws, r, 15, lbl,  bold=True, fill="D9D9D9")
        _cell(ws, r, 16, val,  fill=None)

    # ヘッダー
    ws.cell(37, FREQ_COL, "sweep data")
    _cell(ws, 38, FREQ_COL, "",            bold=True, fill="D9D9D9")
    _cell(ws, 39, FREQ_COL, "Freq [Hz]",   bold=True, fill="D9D9D9")
    for j, e in enumerate(entries):
        _cell(ws, 38, DATA_COL + j, e["level"], bold=True, fill="D9D9D9")
        _cell(ws, 39, DATA_COL + j, "SPL [dB]", bold=True, fill="D9D9D9")
        ws.column_dimensions[
            __import__("openpyxl").utils.get_column_letter(DATA_COL + j)].width = 9

    import openpyxl
    ws.column_dimensions[openpyxl.utils.get_column_letter(FREQ_COL)].width = 11

    # データ
    all_freqs = sorted(set().union(*[set(e["freq_spl"].keys()) for e in entries]))
    data_font = Font(name="Meiryo UI", size=8)
    for i, freq in enumerate(all_freqs):
        row = DATA_START + i
        fc = ws.cell(row, FREQ_COL, freq)
        fc.border = _thin_border(); fc.font = data_font
        fc.number_format = "0.000"
        for j, e in enumerate(entries):
            spl = e["freq_spl"].get(freq)
            sc = ws.cell(row, DATA_COL + j, round(spl, 4) if spl else None)
            sc.border = _thin_border(); sc.font = data_font
            if spl: sc.number_format = "0.00"

    return ws, all_freqs, DATA_START


def _add_chart(wb, ws, group_name, entries, all_freqs, data_start_row):
    """ScatterChart + log X軸をワークシートに追加"""
    n      = len(all_freqs)
    n_lbl  = len(LABEL_FREQS)
    FREQ_COL = 4
    DATA_COL = 5

    # ChartData 隠しシート
    cd_name = f"_CD_{group_name}"
    if cd_name in wb.sheetnames: del wb[cd_name]
    cd = wb.create_sheet(cd_name); cd.sheet_state = "hidden"
    cd.cell(1, 1, "Freq"); cd.cell(1, 2, "SPL")

    # 最初の系列のデータだけチャート用に書く（複数系列対応は後述）
    si = sorted(entries[0]["freq_spl"].items())
    for i, (f, s) in enumerate(si):
        cd.cell(i + 2, 1, f); cd.cell(i + 2, 2, s)
    n_main = len(si)

    # ダミー系列 (カスタムラベル用)
    for i, f in enumerate(LABEL_FREQS):
        cd.cell(i + 2, 3, f)
        cd.cell(i + 2, 4, 21.0)

    chart = ScatterChart()
    chart.scatterStyle = "line"
    chart.title  = "ASA-10 Sine Sweep"
    chart.height = 14; chart.width = 24

    chart.y_axis.scaling.min  = 20; chart.y_axis.scaling.max = 110
    chart.y_axis.majorUnit    = 10
    chart.y_axis.title        = "SPL [dB]"
    chart.y_axis.majorGridlines = ChartLines()
    chart.y_axis.delete       = False

    chart.x_axis.scaling.logBase = 10
    chart.x_axis.scaling.min     = 10
    chart.x_axis.scaling.max     = 20000
    chart.x_axis.title           = "Frequency [Hz]"
    chart.x_axis.majorGridlines  = ChartLines()
    chart.x_axis.minorGridlines  = ChartLines()
    chart.x_axis.delete          = False

    # メイン系列
    x1 = Reference(cd, min_col=1, min_row=2, max_row=n_main + 1)
    y1 = Reference(cd, min_col=2, min_row=1, max_row=n_main + 1)
    s1 = Series(y1, x1, title_from_data=True)
    s1.graphicalProperties.line.solidFill = _PALETTE[0]
    s1.graphicalProperties.line.width     = 15240
    chart.series.append(s1)

    # 追加系列 (複数CSVのとき)
    for k, entry in enumerate(entries[1:], 1):
        si_k = sorted(entry["freq_spl"].items())
        col_x = 5 + (k - 1) * 2
        col_y = 6 + (k - 1) * 2
        for i, (f, s) in enumerate(si_k):
            cd.cell(i + 2, col_x, f); cd.cell(i + 2, col_y, s)
        nk = len(si_k)
        xk = Reference(cd, min_col=col_x, min_row=2, max_row=nk + 1)
        yk = Reference(cd, min_col=col_y, min_row=2, max_row=nk + 1)
        sk = Series(yk, xk)
        sk.graphicalProperties.line.solidFill = _PALETTE[k % len(_PALETTE)]
        sk.graphicalProperties.line.width     = 15240
        chart.series.append(sk)

    # ダミー系列
    x2 = Reference(cd, min_col=3, min_row=2, max_row=n_lbl + 1)
    y2 = Reference(cd, min_col=4, min_row=2, max_row=n_lbl + 1)
    s2 = Series(y2, x2)
    chart.series.append(s2)

    ws.add_chart(chart, "S7")


def patch_xml(x):
    """chart XML に対数グリッド・カスタムラベルなどを適用"""
    # タイトル 11pt
    rpr = '<a:rPr sz="1100" b="0"><a:latin typeface="MS Gothic"/></a:rPr>'
    x = x.replace(
        '<a:r><a:t>ASA-10 Sine Sweep</a:t></a:r>',
        f'<a:r>{rpr}<a:t>ASA-10 Sine Sweep</a:t></a:r>', 1
    )
    # defRPr 9pt
    x = x.replace(
        '<a:defRPr/>',
        '<a:defRPr sz="900" b="0"><a:latin typeface="MS Gothic"/></a:defRPr>'
    )
    # overlay=0
    x = x.replace('</tx></title>', '</tx><overlay val="0"/></title>')
    # axPos バグ修正
    x = x.replace('<axPos val="l"/>', '<axPos val="b"/>', 1)

    # major グリッド D9D9D9
    major_sp = (
        '<majorGridlines><spPr>'
        f'<a:ln xmlns:a="{A}" w="3810">'
        '<a:solidFill><a:srgbClr val="D9D9D9"/></a:solidFill>'
        '</a:ln></spPr></majorGridlines>'
    )
    x = x.replace('<majorGridlines/>', major_sp)

    # minor グリッド EBEBEB (log paper)
    minor_sp = (
        '<minorGridlines><spPr>'
        f'<a:ln xmlns:a="{A}" w="2540">'
        '<a:solidFill><a:srgbClr val="EBEBEB"/></a:solidFill>'
        '</a:ln></spPr></minorGridlines>'
    )
    x = x.replace('<minorGridlines/>', minor_sp)

    # X軸: 白文字28pt でデフォルトラベルを不可視スペーサーに
    txpr_x_white = (
        '<txPr>'
        f'<a:bodyPr xmlns:a="{A}"/>'
        f'<a:lstStyle xmlns:a="{A}"/>'
        f'<a:p xmlns:a="{A}"><a:pPr>'
        '<a:defRPr sz="2800" b="0">'
        f'<a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>'
        '<a:latin typeface="MS Gothic"/>'
        '</a:defRPr>'
        '</a:pPr></a:p>'
        '</txPr>'
    )
    x = re.sub(
        r'(<valAx>(?:(?!</valAx>).)*<axPos val="b"/>)((?:(?!</valAx>).)*)(</valAx>)',
        lambda m: m.group(1) + txpr_x_white + m.group(2) + m.group(3),
        x, count=1, flags=re.DOTALL
    )

    # Y軸 8pt
    txpr_y = (
        '<txPr>'
        f'<a:bodyPr xmlns:a="{A}"/>'
        f'<a:lstStyle xmlns:a="{A}"/>'
        f'<a:p xmlns:a="{A}"><a:pPr>'
        '<a:defRPr sz="800" b="0"><a:latin typeface="MS Gothic"/></a:defRPr>'
        '</a:pPr></a:p></txPr>'
    )
    x = re.sub(
        r'(<valAx>(?:(?!</valAx>).)*<axPos val="l"/>(?:(?!</valAx>).)*)(</valAx>)',
        lambda m: m.group(1) + txpr_y + m.group(2),
        x, count=1, flags=re.DOTALL
    )

    # ダミー系列 (最後の <ser>): カスタムラベル dLbl
    def make_dlbl(idx, text):
        rpr = f'<a:rPr xmlns:a="{A}" sz="700" b="0"><a:latin typeface="MS Gothic"/></a:rPr>'
        return (
            f'<dLbl><idx val="{idx}"/>'
            f'<spPr><a:noFill xmlns:a="{A}"/>'
            f'<a:ln xmlns:a="{A}"><a:noFill/></a:ln></spPr>'
            f'<tx><rich>'
            f'<a:bodyPr xmlns:a="{A}"/>'
            f'<a:lstStyle xmlns:a="{A}"/>'
            f'<a:p xmlns:a="{A}"><a:r>{rpr}<a:t>{text}</a:t></a:r></a:p>'
            f'</rich></tx>'
            f'<showLegendKey val="0"/><showVal val="0"/>'
            f'<showCatName val="0"/><showSerName val="0"/>'
            f'<showPercent val="0"/><showBubbleSize val="0"/>'
            f'<dLblPos val="b"/>'
            f'</dLbl>'
        )

    dlbls_xml = (
        '<dLbls>'
        + ''.join(make_dlbl(i, freq_to_label(f)) for i, f in enumerate(LABEL_FREQS))
        + '<showLegendKey val="0"/><showVal val="0"/>'
        '<showCatName val="0"/><showSerName val="0"/>'
        '<showPercent val="0"/><showBubbleSize val="0"/>'
        '</dLbls>'
    )
    sers = list(re.finditer(r'<ser>.*?</ser>', x, flags=re.DOTALL))
    if sers:
        start, end = sers[-1].start(), sers[-1].end()
        new_ser = re.sub(
            r'<spPr>.*?</spPr>',
            f'<spPr><a:ln xmlns:a="{A}"><a:noFill/></a:ln></spPr>',
            x[start:end], count=1, flags=re.DOTALL
        )
        new_ser = new_ser.replace('<xVal>', dlbls_xml + '<xVal>', 1)
        x = x[:start] + new_ser + x[end:]

    # 凡例削除
    x = re.sub(r'<legend>.*?</legend>', '', x, flags=re.DOTALL)

    # プロットエリア レイアウト
    layout_xml = (
        '<layout><manualLayout>'
        '<layoutTarget val="inner"/>'
        '<xMode val="factor"/><yMode val="factor"/>'
        '<x val="0.06"/><y val="0.04"/>'
        '<w val="0.88"/><h val="0.80"/>'
        '</manualLayout></layout>'
    )
    x = x.replace('<plotArea>', '<plotArea>' + layout_xml, 1)

    # プロットエリア枠線
    plot_sp = (
        '<spPr>'
        f'<a:noFill xmlns:a="{A}"/>'
        f'<a:ln xmlns:a="{A}" w="6350">'
        '<a:solidFill><a:srgbClr val="808080"/></a:solidFill>'
        '</a:ln></spPr>'
    )
    x = x.replace('</plotArea>', plot_sp + '</plotArea>', 1)

    # チャートエリア
    chart_sp = (
        '<spPr>'
        f'<a:solidFill xmlns:a="{A}"><a:srgbClr val="FFFFFF"/></a:solidFill>'
        f'<a:ln xmlns:a="{A}"><a:noFill/></a:ln></spPr>'
    )
    x = x.replace('</chartSpace>', chart_sp + '</chartSpace>', 1)

    return x


def apply_xml_patches(xlsx_path: str):
    """保存済み xlsx の chart XML にパッチ適用"""
    tmp = xlsx_path + ".patching"
    shutil.copy2(xlsx_path, tmp)
    with zipfile.ZipFile(tmp, 'r') as zin:
        with zipfile.ZipFile(xlsx_path, 'w', zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if re.match(r'xl/charts/chart\d+\.xml$', item.filename):
                    data = patch_xml(data.decode('utf-8')).encode('utf-8')
                zout.writestr(item, data)
    os.remove(tmp)


def build_excel(groups: list, settings: dict, output_path: str):
    """グループリストから Excel ファイルを生成"""
    tmp_fd, tmp_path = tempfile.mkstemp(suffix=".xlsx", dir=SCRIPT_DIR)
    os.close(tmp_fd)

    # テンプレートとして空ブックを作成
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)

    chart_targets = []
    for grp in groups:
        ws, all_freqs, data_start = _make_sheet(
            wb, grp["name"], grp["entries"], settings)
        chart_targets.append((grp["name"], grp["entries"], all_freqs, data_start))

    wb.save(tmp_path)

    # チャート追加のため再ロード
    wb2 = load_workbook(tmp_path)
    for gname, entries, all_freqs, data_start in chart_targets:
        ws2 = wb2[gname]
        _add_chart(wb2, ws2, gname, entries, all_freqs, data_start)

    wb2.save(output_path)
    os.remove(tmp_path)

    apply_xml_patches(output_path)


# ── 設定 I/O ─────────────────────────────────────────────────────────────────

def load_settings():
    if os.path.exists(SETTINGS_FILE):
        try:
            with open(SETTINGS_FILE, "r", encoding="utf-8") as f:
                d = json.load(f)
                for k, v in DEFAULT_SETTINGS.items():
                    d.setdefault(k, v)
                return d
        except Exception:
            pass
    return DEFAULT_SETTINGS.copy()


def save_settings(s: dict):
    with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
        json.dump(s, f, ensure_ascii=False, indent=2)


# ── GUI ───────────────────────────────────────────────────────────────────────

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Sine Sweep → Excel 変換 (GitHub版)")
        sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
        w, h   = min(1200, sw - 50), min(780, sh - 50)
        self.geometry(f"{w}x{h}+{(sw-w)//2}+{(sh-h)//2}")
        self.configure(bg=COLOR_BG)
        self._settings = load_settings()
        self._entries: list = []
        self._build_ui()

    def _sec(self, parent, text):
        f = tk.Frame(parent, bg=COLOR_SEC)
        f.pack(fill="x", pady=(10, 2))
        tk.Label(f, text=text, font=("Meiryo UI", 10, "bold"),
                 bg=COLOR_SEC, fg="#333", padx=8, pady=4).pack(side=tk.LEFT)

    def _build_ui(self):
        root = tk.Frame(self, bg=COLOR_BG)
        root.pack(fill="both", expand=True, padx=16, pady=10)
        left  = tk.Frame(root, bg=COLOR_BG); left.pack(side=tk.LEFT, fill="both", expand=True, padx=(0,8))
        right = tk.Frame(root, bg=COLOR_BG); right.pack(side=tk.LEFT, fill="both", expand=True, padx=(8,0))
        self._build_files(left)
        self._build_conditions(right)
        self._build_execute(right)

    def _build_files(self, parent):
        self._sec(parent, "ファイル設定")
        btn_f = tk.Frame(parent, bg=COLOR_BG); btn_f.pack(fill="x", pady=4)
        tk.Button(btn_f, text=" CSV追加 ", bg=COLOR_BTN, fg=COLOR_FG,
                  font=("Meiryo UI", 9), relief="flat", cursor="hand2",
                  command=self._add_csv).pack(side=tk.LEFT, padx=4)
        tk.Button(btn_f, text=" 削除 ", bg="#c0392b", fg=COLOR_FG,
                  font=("Meiryo UI", 9), relief="flat", cursor="hand2",
                  command=self._remove_selected).pack(side=tk.LEFT, padx=4)
        tk.Label(btn_f, text="ダブルクリックでグループ/レベル編集",
                 font=("Meiryo UI", 8), bg=COLOR_BG, fg="#666").pack(side=tk.LEFT, padx=8)

        tf = tk.Frame(parent, bg=COLOR_BG); tf.pack(fill="both", expand=True, pady=4)
        cols = ("filename", "group", "level")
        self._tree = ttk.Treeview(tf, columns=cols, show="headings", height=18)
        self._tree.heading("filename", text="ファイル名")
        self._tree.heading("group",    text="グループ（シート名）")
        self._tree.heading("level",    text="レベル名")
        self._tree.column("filename", width=260); self._tree.column("group", width=180); self._tree.column("level", width=90)
        vsb = ttk.Scrollbar(tf, orient="vertical", command=self._tree.yview)
        self._tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side=tk.RIGHT, fill="y"); self._tree.pack(fill="both", expand=True)
        self._tree.bind("<Double-1>", self._on_double_click)
        self._tree.bind("<Delete>",   lambda e: self._remove_selected())

    def _build_conditions(self, parent):
        self._sec(parent, "測定条件")
        s = self._settings

        def _combo(lbl, key, opts, ck=None):
            f = tk.Frame(parent, bg=COLOR_BG); f.pack(fill="x", padx=8, pady=3)
            tk.Label(f, text=lbl, width=12, anchor="w", font=("Meiryo UI",9), bg=COLOR_BG).pack(side=tk.LEFT)
            var = tk.StringVar(value=s.get(key, ""))
            ttk.Combobox(f, textvariable=var, values=opts, width=25).pack(side=tk.LEFT, fill="x", expand=True)
            setattr(self, f"_v_{key}", var)
            if ck:
                cv = tk.StringVar(value=s.get(ck, ""))
                tk.Entry(f, textvariable=cv, width=20, font=("Meiryo UI",9)).pack(side=tk.LEFT, padx=4)
                setattr(self, f"_v_{ck}", cv)

        def _entry(lbl, key):
            f = tk.Frame(parent, bg=COLOR_BG); f.pack(fill="x", padx=8, pady=3)
            tk.Label(f, text=lbl, width=12, anchor="w", font=("Meiryo UI",9), bg=COLOR_BG).pack(side=tk.LEFT)
            var = tk.StringVar(value=s.get(key, ""))
            tk.Entry(f, textvariable=var, font=("Meiryo UI",9)).pack(side=tk.LEFT, fill="x", expand=True)
            setattr(self, f"_v_{key}", var)

        _combo("Place",   "place",   PLACE_OPTIONS,   "place_custom")
        _combo("Car",     "car",     CAR_OPTIONS,     "car_custom")
        _combo("Signal",  "signal",  SIGNAL_OPTIONS)
        _combo("Mic pos", "mic_pos", MIC_POS_OPTIONS, "mic_pos_custom")
        _combo("Mic",     "mic",     MIC_OPTIONS)
        _entry("Note",    "note")

        self._sec(parent, "出力設定")
        f = tk.Frame(parent, bg=COLOR_BG); f.pack(fill="x", padx=8, pady=3)
        tk.Label(f, text="出力フォルダ", width=12, anchor="w", font=("Meiryo UI",9), bg=COLOR_BG).pack(side=tk.LEFT)
        self._v_output_dir = tk.StringVar(value=s.get("output_dir", ""))
        tk.Entry(f, textvariable=self._v_output_dir, font=("Meiryo UI",9)).pack(side=tk.LEFT, fill="x", expand=True)
        tk.Button(f, text="参照", bg=COLOR_BTN, fg=COLOR_FG, font=("Meiryo UI",8), relief="flat",
                  command=lambda: self._v_output_dir.set(filedialog.askdirectory() or self._v_output_dir.get())
                  ).pack(side=tk.LEFT, padx=4)

        f2 = tk.Frame(parent, bg=COLOR_BG); f2.pack(fill="x", padx=8, pady=3)
        tk.Label(f2, text="ファイル名", width=12, anchor="w", font=("Meiryo UI",9), bg=COLOR_BG).pack(side=tk.LEFT)
        self._v_output_name = tk.StringVar(value=s.get("output_name", "出力結果_Sweep"))
        tk.Entry(f2, textvariable=self._v_output_name, font=("Meiryo UI",9)).pack(side=tk.LEFT, fill="x", expand=True)
        tk.Label(f2, text=".xlsx", font=("Meiryo UI",9), bg=COLOR_BG).pack(side=tk.LEFT)

    def _build_execute(self, parent):
        f = tk.Frame(parent, bg=COLOR_BG); f.pack(fill="x", pady=16)
        self._status = tk.StringVar(value="")
        tk.Label(f, textvariable=self._status, font=("Meiryo UI",9),
                 bg=COLOR_BG, fg="#333", wraplength=400, justify="left").pack(pady=4)
        tk.Button(f, text="  Excel に変換  ", bg=COLOR_RUN, fg=COLOR_FG,
                  font=("Meiryo UI", 12, "bold"), relief="flat", cursor="hand2",
                  command=self._execute).pack(pady=8, ipadx=16, ipady=6)

    def _add_csv(self):
        paths = filedialog.askopenfilenames(
            title="スイープ測定 CSV を選択",
            filetypes=[("CSV", "*.csv;*.CSV"), ("All", "*.*")])
        bad = []
        for path in paths:
            data = parse_sine_csv(path)
            if not data:
                bad.append(os.path.basename(path)); continue
            group, level = extract_group_level(path)
            iid = self._tree.insert("", "end", values=(os.path.basename(path), group, level))
            self._entries.append({"iid": iid, "path": path, "group": group, "level": level})
        if bad:
            messagebox.showwarning("フォーマットエラー",
                "対応フォーマットではないファイル:\n\n" + "\n".join(bad))

    def _remove_selected(self):
        for iid in self._tree.selection():
            self._tree.delete(iid)
            self._entries = [e for e in self._entries if e["iid"] != iid]

    def _on_double_click(self, event):
        if self._tree.identify_region(event.x, event.y) != "cell": return
        col_id = self._tree.identify_column(event.x)
        iid    = self._tree.identify_row(event.y)
        if not iid: return
        mapping = {"#2": ("group", 1), "#3": ("level", 2)}.get(col_id)
        if not mapping: return
        field, ci = mapping
        bbox = self._tree.bbox(iid, col_id)
        if not bbox: return
        x, y, w, h = bbox
        var = tk.StringVar(value=self._tree.item(iid, "values")[ci])
        ent = tk.Entry(self._tree, textvariable=var, font=("Meiryo UI", 9))
        ent.place(x=x, y=y, width=w, height=h); ent.select_range(0, "end"); ent.focus_set()

        def commit(e=None):
            vals = list(self._tree.item(iid, "values"))
            vals[ci] = var.get(); self._tree.item(iid, values=vals)
            for entry in self._entries:
                if entry["iid"] == iid:
                    entry[field] = var.get(); break
            ent.destroy()

        ent.bind("<Return>", commit); ent.bind("<Tab>", commit)
        ent.bind("<Escape>", lambda e: ent.destroy()); ent.bind("<FocusOut>", commit)

    def _collect_settings(self):
        keys = ["place", "place_custom", "car", "car_custom",
                "signal", "mic", "mic_pos", "mic_pos_custom",
                "note", "output_dir", "output_name"]
        return {k: getattr(self, f"_v_{k}", tk.StringVar()).get() for k in keys}

    def _execute(self):
        if not self._entries:
            messagebox.showwarning("入力エラー", "CSVファイルを追加してください。"); return
        s = self._collect_settings(); save_settings(s)

        out_dir  = s.get("output_dir", "").strip() or os.path.dirname(self._entries[0]["path"])
        out_name = (s.get("output_name", "出力結果_Sweep").strip() or "出力結果_Sweep")
        if not out_name.endswith(".xlsx"): out_name += ".xlsx"
        output_path = os.path.join(out_dir, out_name)

        groups_dict: dict = {}
        for e in self._entries:
            groups_dict.setdefault(e["group"], []).append(e)

        self._status.set("CSV読み込み中..."); self.update()

        groups = []
        for gname, elist in groups_dict.items():
            entries = []
            for e in elist:
                freq_spl = parse_sine_csv(e["path"])
                if freq_spl:
                    entries.append({"path": e["path"], "level": e["level"], "freq_spl": freq_spl})
            if entries:
                groups.append({"name": gname, "entries": entries})

        if not groups:
            messagebox.showerror("エラー", "有効なCSVデータがありません。"); return

        self._status.set("Excel生成中..."); self.update()
        try:
            s["date"] = datetime.date.today().isoformat()
            build_excel(groups, s, output_path)
            self._status.set(f"✅ 完了: {output_path}")
            try: os.startfile(output_path)
            except OSError: pass
        except Exception as ex:
            import traceback; traceback.print_exc()
            self._status.set(f"❌ エラー: {ex}")
            messagebox.showerror("エラー", str(ex))


if __name__ == "__main__":
    App().mainloop()
